import openpyxl
import json
import re

def clean_str(val):
    if val is None:
        return None
    s = str(val).strip()
    return s if s else None

def parse_date(val):
    if val is None:
        return None
    s = str(val).strip()
    if not s or s.lower() == 'none':
        return None
    # match YYYY-MM-DD
    m = re.search(r'(\d{4}-\d{2}-\d{2})', s)
    if m:
        return m.group(1)
    return None

def main():
    wb = openpyxl.load_workbook("Project Detail Work Plan RM-Retail POS V2.0_20261015.xlsx", data_only=True)
    ws = wb['Project Plan Updated']

    status_color_map = {
        'Complete': '#77c56b',
        'On Schedule': '#9fc7ff',
        'Late': '#dc3d78',
        'Not Started': '#7b879a'
    }

    milestone_colors = {
        '1': '#ffe86a',
        '2': '#7048c9',
        '3': '#7ce0cc',
        '4': '#ffc8d2'
    }

    project = None
    milestones = []
    current_milestone = None
    all_tasks = []

    task_id_counter = 1
    stage_id_counter = 1

    for r in range(3, ws.max_row + 1):
        wbs_val = ws.cell(row=r, column=2).value
        name_val = ws.cell(row=r, column=3).value
        res_val = ws.cell(row=r, column=4).value
        pct_val = ws.cell(row=r, column=5).value
        pct_work_val = ws.cell(row=r, column=6).value
        status_val = ws.cell(row=r, column=7).value
        dur_val = ws.cell(row=r, column=8).value
        start_val = ws.cell(row=r, column=9).value
        finish_val = ws.cell(row=r, column=10).value
        act_start_val = ws.cell(row=r, column=11).value
        act_finish_val = ws.cell(row=r, column=12).value
        remark_val = ws.cell(row=r, column=13).value

        if wbs_val is None and name_val is None:
            continue

        wbs = clean_str(wbs_val) or ''
        name = clean_str(name_val) or ''
        resources = clean_str(res_val) or ''
        status = clean_str(status_val) or 'Not Started'
        duration = clean_str(dur_val) or ''
        start_date = parse_date(start_val)
        finish_date = parse_date(finish_val)
        act_start = parse_date(act_start_val)
        act_finish = parse_date(act_finish_val)
        remark = clean_str(remark_val)

        pct = round(float(pct_val) * 100, 2) if pct_val is not None else 0.0
        pct_work = round(float(pct_work_val) * 100, 2) if pct_work_val is not None else 0.0

        parts = wbs.split('.')
        level = len(parts)

        # Level 1 / WBS 0 -> Root Project
        if wbs == '0':
            project = {
                'id': 1,
                'name': name,
                'description': f"โครงการพัฒนาระบบ POS ตัวใหม่สำหรับ RM-Retail ({start_date} ถึง {finish_date})",
                'assignee': 'RM, VTEC',
                'status': 'Active',
                'wbs': wbs,
                'start_date': start_date,
                'end_date': finish_date,
                'duration': duration,
                'percent_complete': pct,
                'percent_work_complete': pct_work,
                'project_status': status
            }
            continue

        # Level 1 (WBS 1, 2, 3, 4) -> Milestone (Stage)
        if level == 1 and wbs in ['1', '2', '3', '4']:
            weight_match = re.search(r'\((\d+)%\)', name)
            weight = int(weight_match.group(1)) if weight_match else 25

            stage_obj = {
                'id': stage_id_counter,
                'project_id': 1,
                'wbs': wbs,
                'title': name,
                'start_date': start_date,
                'end_date': finish_date,
                'color': milestone_colors.get(wbs, '#7b879a'),
                'sort_order': stage_id_counter,
                'assignee': resources,
                'duration': duration,
                'weight_pct': weight,
                'percent_complete': pct,
                'percent_work_complete': pct_work,
                'status': status,
                'is_completed': 1 if status == 'Complete' or pct >= 100.0 else 0,
                'tasks_count': 0
            }
            milestones.append(stage_obj)
            current_milestone = stage_obj
            stage_id_counter += 1

            # Milestone 4 has 0 children in Excel, so we also add it as a task under stage 4
            if wbs == '4':
                task_color = status_color_map.get(status, '#7b879a')
                task_obj = {
                    'id': task_id_counter,
                    'stage_id': stage_obj['id'],
                    'wbs': '4.1',
                    'title': name,
                    'level': 2,
                    'parent_wbs': '4',
                    'start_date': start_date,
                    'end_date': finish_date,
                    'duration': duration,
                    'color': task_color,
                    'lane': 1,
                    'assignee': resources,
                    'sort_order': task_id_counter,
                    'percent_complete': pct,
                    'percent_work_complete': pct_work,
                    'status': status,
                    'actual_start': act_start,
                    'actual_finish': act_finish,
                    'is_completed': 1 if status == 'Complete' or pct >= 100.0 else 0,
                    'stuck_reason': 'ล่าช้ากว่ากำหนด' if status == 'Late' else None,
                    'remark': remark
                }
                all_tasks.append(task_obj)
                stage_obj['tasks_count'] += 1
                task_id_counter += 1
            continue

        # Sub-tasks (Level 2, 3, 4)
        if current_milestone:
            parent_wbs = '.'.join(parts[:-1])
            task_color = status_color_map.get(status, '#7b879a')
            is_done = 1 if status == 'Complete' or pct >= 100.0 else 0

            task_obj = {
                'id': task_id_counter,
                'stage_id': current_milestone['id'],
                'wbs': wbs,
                'title': name,
                'level': level,
                'parent_wbs': parent_wbs,
                'start_date': start_date,
                'end_date': finish_date,
                'duration': duration,
                'color': task_color,
                'lane': 1,
                'assignee': resources,
                'sort_order': task_id_counter,
                'percent_complete': pct,
                'percent_work_complete': pct_work,
                'status': status,
                'actual_start': act_start,
                'actual_finish': act_finish,
                'is_completed': is_done,
                'stuck_reason': 'ล่าช้ากว่ากำหนดตามตารางแผนงาน' if status == 'Late' else None,
                'remark': remark
            }
            all_tasks.append(task_obj)
            current_milestone['tasks_count'] += 1
            task_id_counter += 1

    # Resources
    resources_set = set()
    for t in all_tasks:
        if t['assignee']:
            for s in t['assignee'].replace('/', ',').split(','):
                s = s.strip()
                if s:
                    resources_set.add(s)

    assignees_list = sorted(list(resources_set))

    dataset = {
        'project': project,
        'milestones': milestones,
        'tasks': all_tasks,
        'assignees': assignees_list,
        'status_color_map': status_color_map,
        'total_items': len(all_tasks) + len(milestones) + 1
    }

    with open('data/workplan_dataset.json', 'w', encoding='utf-8') as f:
        json.dump(dataset, f, ensure_ascii=False, indent=2)

    print("=== Extraction Summary ===")
    print(f"Project: {project['name']} ({project['start_date']} ~ {project['end_date']})")
    print(f"Milestones ({len(milestones)}):")
    for m in milestones:
        print(f"  - WBS {m['wbs']}: {m['title']} ({m['start_date']} ~ {m['end_date']}) | Tasks: {m['tasks_count']} | Weight: {m['weight_pct']}% | Progress: {m['percent_complete']}%")
    print(f"Total Tasks: {len(all_tasks)}")
    print(f"Assignees ({len(assignees_list)}): {assignees_list}")
    print(f"Saved to data/workplan_dataset.json successfully!")

if __name__ == '__main__':
    main()
